import streamlit as st
import pandas as pd
import plotly.express as px
import re
import math
import ollama
from io import BytesIO
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from src.preprocessing import (
    load_data,
    clean_data
)

from src.feature_engineering import (
    create_features
)

from src.clustering import (
    run_clustering
)

from src.anomaly_detection import (
    run_anomaly_detection
)

from src.insights import (
    create_summary,
    generate_insights
)

# -------------------------------------------------
# AI GRANITE FUNCTION
# -------------------------------------------------
def ask_granite(question, df, summary):

    inventory_context = f"""
You are an AI Inventory Assistant inside a Smart Inventory Manager.

You must answer questions using the inventory information provided below.

Inventory Summary:
- Total Products: {summary['total_products']}
- Total Stock: {summary['total_stock']}
- Total Units Sold: {summary['total_sales']}
- Total Revenue: ₹{summary['total_revenue']:,.2f}

Fast-Moving Products:
{df[df['Movement_Status'] == 'Fast-Moving']
[['Product_Name', 'Stock', 'Units_Sold', 'Revenue']]
.to_string(index=False)}

Slow-Moving Products:
{df[df['Movement_Status'] == 'Slow-Moving']
[['Product_Name', 'Stock', 'Units_Sold', 'Revenue']]
.to_string(index=False)}

Anomalous Products:
{df[df['Anomaly'] == 'Anomaly']
[['Product_Name', 'Stock', 'Units_Sold', 'Revenue', 'Anomaly_Score']]
.to_string(index=False)}

Complete Inventory Data:
{df[
    ['Product_ID',
     'Product_Name',
     'Category',
     'Stock',
     'Units_Sold',
     'Price',
     'Revenue',
     'Movement_Status',
     'Anomaly']
].to_string(index=False)}

Rules:
- Answer only from the provided inventory data.
- Do not invent product information.
- Give concise and useful answers.
- If the requested information is not available, say so.
- When appropriate, provide a practical inventory recommendation.

User Question:
{question}
"""

    response = ollama.chat(
        model="granite3.3:2b",
        messages=[
            {
                "role": "system",
                "content": inventory_context
            },
            {
                "role": "user",
                "content": question
            }
        ]
    )

    return response["message"]["content"]

# -------------------------------------------------
# PAGE CONFIGURATION
# -------------------------------------------------

st.set_page_config(
    page_title="Smart Inventory Manager",
    page_icon="📦",
    layout="wide"
)
st.markdown("""
<style>
    .block-container {
        padding-top: 1.3rem;
    }
</style>
""", unsafe_allow_html=True)

# -------------------------------------------------
# TITLE
# -------------------------------------------------

st.title(
    "📦 Smart Inventory Manager"
)

st.write(
    "AI-powered inventory analysis, "
    "product movement classification "
    "and anomaly detection."
)


# -------------------------------------------------
# SIDEBAR
# -------------------------------------------------

st.sidebar.header(
    "⚙️ Configuration"
)


uploaded_file = st.sidebar.file_uploader(
    "Upload Inventory File",
    type=[
        "csv",
        "xlsx"
    ]
)


# Number of clusters
number_of_clusters = (
    st.sidebar.slider(
        "Number of Clusters",
        min_value=2,
        max_value=6,
        value=3
    )
)


# Anomaly sensitivity
anomaly_sensitivity = (
    st.sidebar.slider(
        "Anomaly Sensitivity",
        min_value=1,
        max_value=49,
        value=10
    )
)



# -------------------------------------------------
# INPUT COLUMN NORMALIZATION
# -------------------------------------------------

def normalize_input_columns(dataframe):
    dataframe = dataframe.copy()

    aliases = {
        "productid": "Product_ID",
        "product_id": "Product_ID",
        "product name": "Product_Name",
        "productname": "Product_Name",
        "product_name": "Product_Name",
        "category": "Category",
        "stock": "Stock",
        "inventory": "Stock",
        "units sold": "Units_Sold",
        "unitssold": "Units_Sold",
        "units_sold": "Units_Sold",
        "sales": "Units_Sold",
        "quantity sold": "Units_Sold",
        "quantity_sold": "Units_Sold",
        "price": "Price",
        "unit price": "Price",
        "unit_price": "Price",
        "date": "Date",
        "supplier": "Supplier",
        "vendor": "Supplier",
    }

    renamed = {}
    for column in dataframe.columns:
        original = str(column).strip()
        normalized = re.sub(r"[\s\-]+", "_", original).lower()
        compact = normalized.replace("_", "")
        if normalized in aliases:
            renamed[column] = aliases[normalized]
        elif compact in aliases:
            renamed[column] = aliases[compact]
        else:
            renamed[column] = original

    dataframe = dataframe.rename(columns=renamed)
    dataframe = dataframe.loc[:, ~dataframe.columns.duplicated()]
    return dataframe


# -------------------------------------------------
# DATA LOADING
# -------------------------------------------------

try:

    if uploaded_file is not None:

        # Uploaded dataset has priority
        df = load_data(
            uploaded_file
        )

        st.sidebar.success(
            f"Using uploaded file: "
            f"{uploaded_file.name}"
        )

    else:

        # Use sample dataset
        df = load_data(
            "data/inventory.csv"
        )

        st.sidebar.info(
            "Using sample inventory dataset"
        )

    # Automatically standardize common column-name variations.
    df = normalize_input_columns(df)

except Exception as e:

    st.error(
        f"Error loading data: {e}"
    )

    st.stop()


# -------------------------------------------------
# DATA CLEANING
# -------------------------------------------------

try:

    df = clean_data(df)

except Exception as e:

    st.error(
        f"Data validation error: {e}"
    )

    st.info(
        "Please make sure your inventory "
        "file contains these columns:"
    )

    st.code(
        """
Product_ID
Product_Name
Category
Stock
Units_Sold
Price
Date
Supplier
        """
    )

    st.stop()


# -------------------------------------------------
# FEATURE ENGINEERING
# -------------------------------------------------

df = create_features(df)


# -------------------------------------------------
# K-MEANS CLUSTERING
# -------------------------------------------------

df, clustering_model, scaler = (
    run_clustering(
        df,
        number_of_clusters
    )
)


# -------------------------------------------------
# ISOLATION FOREST
# -------------------------------------------------

df, anomaly_model = (
    run_anomaly_detection(
        df,
        anomaly_sensitivity / 100
    )
)


# -------------------------------------------------
# SUMMARY
# -------------------------------------------------

summary = create_summary(
    df
)

# =================================================
# AI INVENTORY ASSISTANT
# =================================================

def inventory_ai_response(question, df, summary):
    question = question.lower().strip()

    # ---------------------------------------------
    # BASIC DATA
    # ---------------------------------------------

    total_products = summary["total_products"]
    total_stock = summary["total_stock"]
    total_sales = summary["total_sales"]
    total_revenue = summary["total_revenue"]

    fast_products = df[
        df["Movement_Status"] == "Fast-Moving"
    ]

    slow_products = df[
        df["Movement_Status"] == "Slow-Moving"
    ]

    normal_products = df[
        df["Movement_Status"] == "Normal-Moving"
    ]

    anomalies = df[
        df["Anomaly"] == "Anomaly"
    ]

    # ---------------------------------------------
    # TOTAL INVENTORY
    # ---------------------------------------------

    if (
        "total revenue" in question
        or "revenue" in question
        and "product" not in question
    ):
        return (
            f"💰 **Total Revenue:** ₹{total_revenue:,.2f}\n\n"
            f"Based on the uploaded inventory dataset, "
            f"the total calculated revenue is ₹{total_revenue:,.2f}."
        )

    if (
        "total stock" in question
        or "how much stock" in question
        or "inventory stock" in question
    ):
        return (
            f"📦 **Total Stock:** {total_stock:,} units\n\n"
            f"The uploaded dataset currently contains "
            f"{total_stock:,} units of inventory."
        )

    if (
        "total sales" in question
        or "units sold" in question
        or "how many sold" in question
    ):
        return (
            f"📈 **Total Units Sold:** {total_sales:,}\n\n"
            f"The inventory dataset contains "
            f"{total_sales:,} units sold."
        )

    if (
        "how many products" in question
        or "number of products" in question
        or "total products" in question
    ):
        return (
            f"📦 **Total Products:** {total_products}\n\n"
            f"The uploaded dataset contains "
            f"{total_products} unique products."
        )

    # ---------------------------------------------
    # FAST MOVING PRODUCTS
    # ---------------------------------------------

    if (
        "fast moving" in question
        or "fast-moving" in question
        or "best selling" in question
        or "best-selling" in question
        or "high sales" in question
    ):

        if fast_products.empty:
            return "No Fast-Moving products were identified in the current analysis."

        top_products = (
            fast_products
            .sort_values("Units_Sold", ascending=False)
            .head(5)
        )

        response = (
            "🚀 **Fast-Moving Products**\n\n"
            "The following products are classified as Fast-Moving "
            "by the K-Means inventory analysis:\n\n"
        )

        for _, row in top_products.iterrows():
            response += (
                f"• **{row['Product_Name']}** — "
                f"{int(row['Units_Sold']):,} units sold, "
                f"Stock: {int(row['Stock']):,}\n"
            )

        return response

    # ---------------------------------------------
    # SLOW MOVING PRODUCTS
    # ---------------------------------------------

    if (
        "slow moving" in question
        or "slow-moving" in question
        or "poor sales" in question
        or "low sales" in question
    ):

        if slow_products.empty:
            return "No Slow-Moving products were identified."

        top_products = (
            slow_products
            .sort_values("Units_Sold", ascending=True)
            .head(5)
        )

        response = (
            "🐢 **Slow-Moving Products**\n\n"
            "The following products have relatively low sales movement:\n\n"
        )

        for _, row in top_products.iterrows():
            response += (
                f"• **{row['Product_Name']}** — "
                f"{int(row['Units_Sold']):,} units sold, "
                f"Stock: {int(row['Stock']):,}\n"
            )

        return response

    # ---------------------------------------------
    # ANOMALIES
    # ---------------------------------------------

    if (
        "anomal" in question
        or "unusual" in question
        or "abnormal" in question
        or "irregular" in question
    ):

        if anomalies.empty:
            return (
                "✅ **No unusual inventory activity detected.**\n\n"
                "Isolation Forest did not identify any products "
                "as anomalies with the current sensitivity setting."
            )

        anomaly_products = anomalies.head(10)

        response = (
            f"🚨 **Inventory Anomalies Detected: {len(anomalies)}**\n\n"
            "Isolation Forest identified the following unusual records:\n\n"
        )

        for _, row in anomaly_products.iterrows():
            response += (
                f"• **{row['Product_Name']}** — "
                f"Stock: {int(row['Stock']):,}, "
                f"Units Sold: {int(row['Units_Sold']):,}, "
                f"Anomaly Score: {row['Anomaly_Score']:.4f}\n"
            )

        return response

    # ---------------------------------------------
    # CATEGORY WITH HIGHEST SALES
    # ---------------------------------------------

    if (
        "highest sales category" in question
        or "best category" in question
        or "highest selling category" in question
        or "which category" in question
    ):

        category_sales = (
            df.groupby("Category")["Units_Sold"]
            .sum()
            .sort_values(ascending=False)
        )

        if not category_sales.empty:
            category = category_sales.index[0]
            sales = int(category_sales.iloc[0])

            return (
                f"🏆 **Top Sales Category:** {category}\n\n"
                f"This category has the highest total unit sales "
                f"with **{sales:,} units sold**."
            )

    # ---------------------------------------------
    # PRODUCT-SPECIFIC SEARCH
    # ---------------------------------------------

    for product in df["Product_Name"].unique():

        if product.lower() in question:

            product_data = df[
                df["Product_Name"] == product
            ]

            if not product_data.empty:

                row = product_data.iloc[0]

                movement = row["Movement_Status"]
                anomaly = row["Anomaly"]

                response = (
                    f"📦 **{row['Product_Name']}**\n\n"
                    f"• Category: **{row['Category']}**\n"
                    f"• Stock: **{int(row['Stock']):,} units**\n"
                    f"• Units Sold: **{int(row['Units_Sold']):,}**\n"
                    f"• Price: **₹{row['Price']:,.2f}**\n"
                    f"• Revenue: **₹{row['Revenue']:,.2f}**\n"
                    f"• Movement: **{movement}**\n"
                    f"• Status: **{anomaly}**\n"
                )

                if movement == "Fast-Moving":
                    response += (
                        "\n💡 **AI Recommendation:** "
                        "This product has strong sales movement. "
                        "Monitor its stock level regularly to avoid "
                        "potential stock shortages."
                    )

                elif movement == "Slow-Moving":
                    response += (
                        "\n💡 **AI Recommendation:** "
                        "This product has relatively low sales movement. "
                        "Review its demand and current stock before "
                        "ordering additional units."
                    )

                else:
                    response += (
                        "\n💡 **AI Recommendation:** "
                        "The product currently shows normal inventory movement."
                    )

                if anomaly == "Anomaly":
                    response += (
                        "\n\n🚨 **Attention:** "
                        "Isolation Forest has also flagged this product "
                        "as having unusual inventory activity."
                    )

                return response

    # ---------------------------------------------
    # RESTOCK / LOW STOCK
    # ---------------------------------------------

    if (
        "restock" in question
        or "replenish" in question
        or "low stock" in question
        or "stock attention" in question
    ):

        median_stock = df["Stock"].median()

        attention = df[
            df["Stock"] < median_stock
        ].sort_values("Units_Sold", ascending=False).head(5)

        if attention.empty:
            return "No products currently require special stock attention based on the available data."

        response = (
            "📦 **Products That May Need Stock Attention**\n\n"
            "These products have relatively low stock compared with "
            "the dataset's median stock level:\n\n"
        )

        for _, row in attention.iterrows():
            response += (
                f"• **{row['Product_Name']}** — "
                f"Stock: {int(row['Stock']):,}, "
                f"Units Sold: {int(row['Units_Sold']):,}\n"
            )

        response += (
            "\n💡 These are data-driven indicators, not automatic "
            "purchase orders. Actual replenishment should also consider "
            "future demand and supplier lead time."
        )

        return response

    # ---------------------------------------------
    # GENERAL INVENTORY SUMMARY
    # ---------------------------------------------

    if (
        "summary" in question
        or "overview" in question
        or "inventory status" in question
        or "analyze my inventory" in question
        or "how is my inventory" in question
    ):

        return (
            "📊 **Inventory Overview**\n\n"
            f"• Products: **{total_products}**\n"
            f"• Total Stock: **{total_stock:,} units**\n"
            f"• Units Sold: **{total_sales:,}**\n"
            f"• Revenue: **₹{total_revenue:,.2f}**\n"
            f"• Fast-Moving: **{len(fast_products)}**\n"
            f"• Normal-Moving: **{len(normal_products)}**\n"
            f"• Slow-Moving: **{len(slow_products)}**\n"
            f"• Anomalies: **{len(anomalies)}**"
        )

    # ---------------------------------------------
    # DEFAULT RESPONSE
    # ---------------------------------------------

    return (
        "🤖 I can help you analyze your inventory data.\n\n"
        "Try asking:\n"
        "• Which products are fast-moving?\n"
        "• Which products are slow-moving?\n"
        "• Which products have unusual activity?\n"
        "• Which category has the highest sales?\n"
        "• Which products may need restocking?\n"
        "• What is my total revenue?\n"
        "• Tell me about Smart TV.\n"
        "• Give me an inventory summary."
    )

# -------------------------------------------------
# TABS
# -------------------------------------------------

tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs(
    [
        "📊 Dashboard",
        "📦 Product Analysis",
        "🚨 Anomaly Detection",
        "📋 Inventory Data",
        "🤖 AI Inventory",
        "⬇️ Download"
    ]
)


# =================================================
# TAB 1 - DASHBOARD
# =================================================

with tab1:

    st.subheader(
        "Product Movement"
    )

    movement_counts = (
        df[
            "Movement_Status"
        ]
        .value_counts()
        .reset_index()
    )

    movement_counts.columns = [
        "Movement_Status",
        "Count"
    ]

    fig = px.bar(
        movement_counts,
        x="Movement_Status",
        y="Count",
        title="Product Movement Classification"
    )

    st.plotly_chart(
        fig,
        use_container_width=True
    )


    st.subheader(
        "Sales by Category"
    )

    category_sales = (
        df.groupby("Category")
        ["Units_Sold"]
        .sum()
        .reset_index()
    )

    fig2 = px.pie(
        category_sales,
        names="Category",
        values="Units_Sold",
        title="Units Sold by Category"
    )

    st.plotly_chart(
        fig2,
        use_container_width=True
    )


    # Insights
    st.subheader(
        "💡 Smart Insights"
    )

    insights = generate_insights(
        df
    )

    if insights:

        for insight in insights:

            st.info(
                insight
            )

    else:

        st.info(
            "No additional insights available."
        )


# =================================================
# TAB 2 - PRODUCT ANALYSIS
# =================================================

with tab2:

    st.subheader(
        "📦 Product Analysis"
    )

    product_names = sorted(
        df["Product_Name"]
        .unique()
        .tolist()
    )

    selected_product = st.selectbox(
        "Select Product",
        product_names
    )

    product_data = df[
        df["Product_Name"]
        == selected_product
    ]

    st.dataframe(
        product_data,
        use_container_width=True
    )

    if not product_data.empty:

        row = product_data.iloc[0]

        c1, c2, c3, c4 = st.columns(4)

        with c1:

            st.metric(
                "Stock",
                int(row["Stock"])
            )

        with c2:

            st.metric(
                "Units Sold",
                int(row["Units_Sold"])
            )

        with c3:

            st.metric(
                "Revenue",
                f"₹{row['Revenue']:,.2f}"
            )

        with c4:

            st.metric(
                "Movement",
                row["Movement_Status"]
            )


# =================================================
# TAB 3 - ANOMALY DETECTION
# =================================================

with tab3:

    st.subheader(
        "🚨 Inventory Anomaly Detection"
    )

    anomaly_data = df[
        df["Anomaly"]
        == "Anomaly"
    ]

    st.metric(
        "Detected Anomalies",
        len(anomaly_data)
    )

    if anomaly_data.empty:

        st.success(
            "No unusual inventory activity detected."
        )

    else:

        st.warning(
            "Unusual inventory activity detected."
        )

        st.dataframe(
            anomaly_data[
                [
                    "Product_ID",
                    "Product_Name",
                    "Category",
                    "Stock",
                    "Units_Sold",
                    "Revenue",
                    "Anomaly_Score"
                ]
            ],
            use_container_width=True
        )


# =================================================
# TAB 4 - INVENTORY DATA
# =================================================

with tab4:

    st.subheader(
        "📋 Inventory Data"
    )

    search = st.text_input(
        "🔎 Search Product"
    )

    filtered_df = df.copy()

    if search:

        mask = (
            filtered_df[
                "Product_Name"
            ]
            .str.contains(
                search,
                case=False,
                na=False
            )
            |
            filtered_df[
                "Product_ID"
            ]
            .str.contains(
                search,
                case=False,
                na=False
            )
        )

        filtered_df = (
            filtered_df[mask]
        )

    st.dataframe(
        filtered_df,
        use_container_width=True
    )

# =================================================
# TAB 5 - AI INVENTORY INTELLIGENCE
# =================================================

with tab5:

    st.subheader(
        "🤖 AI Inventory Intelligence"
    )

    st.write(
        "Use your inventory data and machine-learning results "
        "to get data-driven insights and recommendations."
    )

    # -------------------------------------------------
    # DATA-DRIVEN INVENTORY ADVISOR
    # -------------------------------------------------

    st.markdown("---")

    st.subheader(
        "📊 Data-Driven Inventory Advisor"
    )

    st.caption(
        "Automatically generated from the current inventory analysis."
    )

    fast_products = df[
        df["Movement_Status"] == "Fast-Moving"
    ]

    slow_products = df[
        df["Movement_Status"] == "Slow-Moving"
    ]

    anomalies = df[
        df["Anomaly"] == "Anomaly"
    ]

    # Determine the main finding
    if not anomalies.empty:

        finding_product = anomalies.iloc[0]["Product_Name"]

        finding_text = (
            f"Unusual inventory activity was detected for "
            f"**{finding_product}** based on the Isolation Forest analysis."
        )

        recommendation = (
            "Review its stock level, units sold and recent inventory "
            "activity before making the next inventory decision."
        )

    elif not fast_products.empty:

        finding_product = (
            fast_products
            .sort_values(
                "Units_Sold",
                ascending=False
            )
            .iloc[0]
        )

        finding_text = (
            f"**{finding_product['Product_Name']}** is currently one of "
            f"the strongest fast-moving products based on sales activity."
        )

        recommendation = (
            "Monitor its available stock regularly because continued "
            "high sales activity may increase replenishment requirements."
        )

    elif not slow_products.empty:

        finding_product = (
            slow_products
            .sort_values(
                "Units_Sold",
                ascending=True
            )
            .iloc[0]
        )

        finding_text = (
            f"**{finding_product['Product_Name']}** has relatively low "
            f"sales movement compared with other products."
        )

        recommendation = (
            "Review demand and current stock before placing additional "
            "orders for this product."
        )

    else:

        finding_text = (
            "The current inventory data does not show a major "
            "priority requiring special attention."
        )

        recommendation = (
            "Continue monitoring product movement and inventory activity."
        )

    # Finding box
    st.success(
        f"🎯 **AI Finding**\n\n{finding_text}"
    )

    # Recommendation box
    st.info(
        f"💡 **Recommended Action**\n\n{recommendation}"
    )

    # -------------------------------------------------
    # INVENTORY KNOWLEDGE ADVISOR
    # -------------------------------------------------

    st.markdown("---")

    st.subheader(
        "🔎 Inventory Knowledge Advisor"
    )

    st.caption(
        "Select a question to get an instant data-driven answer."
    )

    advisor_questions = [
        "Which products are fast-moving?",
        "Which products are slow-moving?",
        "Which products have unusual activity?",
        "Which category has the highest sales?",
        "Which products may need restocking?",
        "What is my total revenue?",
        "Give me an inventory summary."
    ]

    selected_question = st.selectbox(
        "Choose a question",
        advisor_questions
    )

    if st.button(
        "🔍 Analyze",
        use_container_width=False
    ):

        answer = inventory_ai_response(
            selected_question,
            df,
            summary
        )

        st.markdown(
            "### 🤖 AI Response"
        )

        st.markdown(
            answer
        )

    # -------------------------------------------------
    # CONVERSATIONAL AI ASSISTANT
    # -------------------------------------------------

    st.markdown("---")

    st.subheader(
        "💬 Conversational AI Inventory Assistant"
    )

    st.caption(
        "Ask questions about your uploaded inventory dataset "
        "in natural language."
    )

    # Initialize chat history
    if "inventory_chat_history" not in st.session_state:

        st.session_state.inventory_chat_history = [
            {
                "role": "assistant",
                "content": (
                    "Hello! 👋 I'm your AI Inventory Assistant. "
                    "Ask me about products, sales, stock, revenue, "
                    "movement or anomalies."
                )
            }
        ]

    # Display previous messages
    for message in st.session_state.inventory_chat_history:

        with st.chat_message(
            message["role"]
        ):

            st.markdown(
                message["content"]
            )

    # Chat input
    user_question = st.chat_input(
        "Ask your inventory question..."
    )

    if user_question:

        # Add user message
        st.session_state.inventory_chat_history.append(
            {
                "role": "user",
                "content": user_question
            }
        )

        # Generate response

        try:

            answer = ask_granite(
        user_question,
        df,
        summary
    )

        except Exception as e:

          answer = (
        "⚠️ **Unable to generate AI response.**\n\n"
        f"**Error:** `{str(e)}`\n\n"
          )

        # Add assistant response
        st.session_state.inventory_chat_history.append(
            {
                "role": "assistant",
                "content": answer
            }
        )

        # Rerun to display messages
        st.rerun()

# =================================================
# TAB 6 - DOWNLOAD
# =================================================
with tab6:
    st.subheader("⬇️ Download Analysis")

    export_df = df.copy()

    if "Anomaly" in export_df.columns:
        export_df = export_df.rename(
            columns={"Anomaly": "Anomaly_Status"}
        )

    if "Anomaly_Status" in export_df.columns:
        export_df["Anomaly_Status"] = export_df[
            "Anomaly_Status"
        ].map(
            {
                "Normal": "Normal",
                "Anomaly": "Anomaly"
            }
        ).fillna("Normal")

    if "Date" in export_df.columns:
        export_df["Date"] = pd.to_datetime(
            export_df["Date"],
            errors="coerce"
        ).dt.strftime("%d-%m-%Y")


    csv_data = export_df.to_csv(
        index=False
    ).encode("utf-8-sig")

    st.download_button(
        label="📄 Download CSV",
        data=csv_data,
        file_name="inventory_analysis.csv",
        mime="text/csv"
    )

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        export_df.to_excel(
            writer,
            index=False,
            sheet_name="Inventory Analysis"
        )

        worksheet = writer.sheets["Inventory Analysis"]

        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        thin_border = Border(left=Side(style="thin", color="B7B7B7"), right=Side(style="thin", color="B7B7B7"), top=Side(style="thin", color="B7B7B7"), bottom=Side(style="thin", color="B7B7B7"))

        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill
            cell.border = thin_border
        worksheet.row_dimensions[1].height = 32

        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = len(str(column_cells[0].value or ""))
            for cell in column_cells[1:]:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 30)

        special_widths = {"Product_ID":14,"Product_Name":22,"Category":18,"Stock":12,"Units_Sold":14,"Price":14,"Date":14,"Supplier":22,"Revenue":16,"Sales_Frequency":20,"Stock_Utilization":20,"Stock_Turnover":18,"Cluster":12,"Movement_Status":20,"Anomaly_Status":18,"Anomaly_Score":18}
        for column_name, width in special_widths.items():
            if column_name in export_df.columns:
                column_letter = get_column_letter(export_df.columns.get_loc(column_name) + 1)
                worksheet.column_dimensions[column_letter].width = width

        if "Date" in export_df.columns:
            date_column_letter = get_column_letter(export_df.columns.get_loc("Date") + 1)
            worksheet.column_dimensions[date_column_letter].width = 14
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet[f"{date_column_letter}{row}"]
                if cell.value is not None:
                    cell.value = str(cell.value)
                cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in range(2, worksheet.max_row + 1):
            max_lines = 1
            for cell in worksheet[row]:
                if cell.value is not None:
                    width = worksheet.column_dimensions[get_column_letter(cell.column)].width or 10
                    lines = max(1, math.ceil(len(str(cell.value)) / max(width - 2, 1)))
                    max_lines = max(max_lines, lines)
            worksheet.row_dimensions[row].height = min(max(18, max_lines * 15), 45)

    output.seek(0)

    st.download_button(
        label="📊 Download Formatted Excel",
        data=output.getvalue(),
        file_name="inventory_analysis.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="download_formatted_excel"
    )
